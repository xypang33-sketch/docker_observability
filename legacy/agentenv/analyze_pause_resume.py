"""Analyze host resource usage during pause/resume windows of replay run(s).

Pairs a high-frequency spotbox_monitor capture (e.g. 0.1s interval) with
per-step pause/resume timestamps, assigns every sample to a phase
(pause / paused / resume / exec / gap), and reports per-phase host
CPU / FC cores / disk IO / dirty-page stats, plus per-cycle peak disk write
(during pause) and peak disk read (during resume).

Timestamp sources:
  default   replay-result.json — harbor job dir layout
            (<job_dir>/*/agent/replay-result.json) or a replay-aenv output
            dir with the file at its top level (single task)
  --runlog  ReplayAgent logs, three formats auto-detected per file:
            - current (main @ec5c928): self-contained
              "Agent interaction timing step=i/N pause=.. delay=.. ..
              lines; anchored at the line's own timestamp (job.log) or the
              step's COMMAND END block timestamp (console/run.log)
            - jobs.txt (older standalone agent): "Lifecycle
              scheduler_wait/execution" lines carry only durations; absolute
              times are reconstructed by anchoring at the same-millisecond
              probe line and subtracting durations along prev RESULT ->
              sched.pause -> pause -> delay -> sched.resume -> resume(+retry)
              -> probe -> COMMAND -> RESULT
            - legacy run.log (harbor replay agent): "Replay lifecycle:
              started_at=... finished_at=..." lines carry absolute timestamps
          Path forms:
            - file                          -> single task (either format)
            - <dir> contains run.log        -> single-task record
            - <dir>/<task&#42;>/run.log         -> multi-task record
            - <dir>/&#42;.txt|&#42;.log (flat)      -> one new-format task per file

Usage: python3 analyze_pause_resume.py <obs_dir> <job_dir> [--plot]
       python3 analyze_pause_resume.py <obs_dir> --runlog <file|dir> [--plot]

--plot writes pause_resume_cycle.png (single: median cycle, zoomed) or
pause_resume_multi.png (multi: per-task phase gantt + host curves).
"""

import glob
import json
import os
import re
import sys
from datetime import datetime, timedelta

PHASES = ("pause", "paused", "resume", "exec", "gap")

LIFECYCLE_RE = re.compile(
    r"Replay lifecycle: started_at=(\S+) (?:finished_at|completed_at)=(\S+) "
    r"delay=([\d.]+)s pause=([\d.]+)s paused=([\d.]+)s resume=([\d.]+)s"
)
CMD_RE = re.compile(r"\[(\S+)\] \[(\d+)/(\d+)\] COMMAND (START|END)")


def to_epoch(ts):
    if not ts:
        return None
    dt = datetime.fromisoformat(ts.replace("Z", "+00:00"))
    if dt.tzinfo is not None:
        dt = dt.astimezone().replace(tzinfo=None)
    return dt.timestamp()


def load_metrics(obs_dir):
    rows = []
    try:
        with open(os.path.join(obs_dir, "metrics_1s.jsonl")) as f:
            for line in f:
                line = line.strip()
                if line:
                    rows.append(json.loads(line))
    except FileNotFoundError:
        pass
    return rows


def load_steps(job_dir):
    direct = os.path.join(job_dir, "replay-result.json")
    if os.path.isfile(direct):
        # replay-aenv output dir (this repo's agent/replay_agent.py): the file
        # sits at the top level and carries the same per-step fields, including
        # absolute pause/resume timestamps.
        hits = [direct]
    else:
        hits = glob.glob(os.path.join(job_dir, "*", "agent", "replay-result.json"))
    if not hits:
        sys.exit(f"no replay-result.json under {job_dir}")
    return json.load(open(hits[0])).get("steps", [])


def _iso(dt):
    return dt.isoformat(timespec="milliseconds")


def parse_runlog(path):
    """Parse one ReplayAgent run.log into replay-result.json-style step dicts.

    Each "Replay lifecycle" line carries started_at (pause begin) /
    finished_at (resume done) plus pause/paused/resume durations; the k-th
    lifecycle line pairs with the k-th COMMAND START/END block (exec window,
    action text, exit_code, elapsed).
    """
    lifecycles = []
    cmds = {}
    cur = None       # idx whose action lines are being read
    last_end = None  # idx whose exit_code/elapsed lines may follow
    with open(path, errors="replace") as f:
        for line in f:
            m = LIFECYCLE_RE.search(line)
            if m:
                lifecycles.append(
                    (datetime.fromisoformat(m.group(1)),
                     datetime.fromisoformat(m.group(2)),
                     float(m.group(4)), float(m.group(5)), float(m.group(6)))
                )
                cur = last_end = None
                continue
            c = CMD_RE.search(line)
            if c:
                idx = int(c.group(2)) - 1
                ts = datetime.fromisoformat(c.group(1))
                if c.group(4) == "START":
                    cmds.setdefault(idx, {})["START"] = ts
                    cmds[idx]["action_lines"] = []
                    cur, last_end = idx, None
                else:
                    cmds.setdefault(idx, {})["END"] = ts
                    cur, last_end = None, idx
                continue
            if cur is not None:
                if line.strip():
                    cmds[cur]["action_lines"].append(line.rstrip("\n"))
                else:
                    cur = None
                continue
            if last_end is not None:
                if line.startswith("exit_code:"):
                    v = line.split(":", 1)[1].strip()
                    cmds[last_end]["exit"] = None if v == "None" else int(v)
                elif line.startswith("elapsed:"):
                    cmds[last_end]["elapsed"] = float(
                        line.split(":", 1)[1].strip().rstrip("s"))
                elif line.startswith("stdout:"):
                    last_end = None
    steps = []
    for i, (started, finished, pause_s, paused_s, resume_s) in enumerate(lifecycles):
        pc = started + timedelta(seconds=pause_s)
        rs = finished - timedelta(seconds=resume_s)
        cmd = cmds.get(i) or {}
        steps.append({
            "index": i,
            "action": " ".join(cmd.get("action_lines") or []),
            "exit": cmd.get("exit"),
            "pause_ms": round(pause_s * 1000, 1),
            "paused_ms": round(paused_s * 1000, 1),
            "resume_ms": round(resume_s * 1000, 1),
            "exec_ms": round((cmd.get("elapsed") or
                              (cmd.get("END", finished) - cmd.get("START", finished))
                              .total_seconds()) * 1000, 1),
            "pause_started_at": _iso(started),
            "pause_completed_at": _iso(pc),
            "resume_started_at": _iso(rs),
            "resume_completed_at": _iso(finished),
            "started_at": _iso(cmd.get("START", finished)),
            "completed_at": _iso(cmd.get("END", finished)),
        })
    return steps


V2_STARTED_RE = re.compile(r"Replay started trajectory=(\S+)")
V2_PROBE_RE = re.compile(
    r"Sandbox ready probe succeeded attempt=\S+ elapsed=([\d.]+)s "
    r"retry_wait=([\d.]+)s")
V2_SCHED_RE = re.compile(
    r"Lifecycle scheduler_wait pause=([\d.]+)s resume=([\d.]+)s")
V2_EXEC_RE = re.compile(
    r"Lifecycle execution pause=([\d.]+)s delay=([\d.]+)s resume=([\d.]+)s "
    r"resume_retry_wait=([\d.]+)s .*pause_error=(\S+)")
V2_CMD_RE = re.compile(r"\[(\d+)/\d+\] COMMAND (.*)")
V2_RESULT_RE = re.compile(
    r"\[(\d+)/\d+\] RESULT exit_code=(-?\d+) succeeded=\w+ elapsed=([\d.]+)s")


def parse_joblog(path):
    """Parse a new-format ReplayAgent log (jobs.txt style) -> [(name, steps)].

    Per step the agent logs a same-millisecond group: "Sandbox ready probe
    succeeded", "Lifecycle scheduler_wait", "Lifecycle execution" (durations
    only, no absolute start) and "[i/N] COMMAND"; "[i/N] RESULT" follows with
    the exec elapsed. Absolute pause/resume boundaries are reconstructed by
    subtracting durations from the group timestamp:
        resume_completed = T_group - probe_elapsed - probe_retry_wait
        resume_started   = resume_completed - resume - resume_retry_wait
        pause_completed  = resume_started - scheduler_wait.resume - delay
        pause_started    = pause_completed - pause
    Validated on real logs: reconstructed pause_started lands ms-exactly on
    the previous step's RESULT timestamp. Segments split on "Replay started"
    (one trajectory each); name comes from the trajectory file stem.
    """
    def parse_ts(s):
        return datetime.fromisoformat(s)

    segments = []  # (name|None, raw: idx -> dict)
    name = None
    raw = None
    pend_probe = pend_sched = pend_exec = None
    with open(path, errors="replace") as f:
        for line in f:
            if not line.startswith("2"):  # fast path: all lines start with ISO ts
                continue
            ts_s, _, msg = line.partition(" INFO ")
            if not msg:
                continue
            ts = parse_ts(ts_s.strip())
            m = V2_STARTED_RE.search(msg)
            if m:
                if raw:
                    segments.append((name, raw))
                name = os.path.splitext(os.path.basename(m.group(1)))[0]
                raw = {}
                pend_probe = pend_sched = pend_exec = None
                continue
            if raw is None:
                continue
            m = V2_PROBE_RE.search(msg)
            if m:
                pend_probe = (ts, float(m.group(1)), float(m.group(2)))
                continue
            m = V2_SCHED_RE.search(msg)
            if m:
                pend_sched = (float(m.group(1)), float(m.group(2)))
                continue
            m = V2_EXEC_RE.search(msg)
            if m:
                pend_exec = (ts, float(m.group(1)), float(m.group(2)),
                             float(m.group(3)), float(m.group(4)),
                             m.group(5))
                continue
            m = V2_CMD_RE.search(msg)
            if m and pend_exec:
                idx = int(m.group(1)) - 1
                raw[idx] = {"lc": pend_exec, "probe": pend_probe,
                            "sched": pend_sched, "cmd_ts": ts,
                            "action": m.group(2).replace("\\n", " ")}
                pend_probe = pend_sched = pend_exec = None
                continue
            m = V2_RESULT_RE.search(msg)
            if m:
                idx = int(m.group(1)) - 1
                raw.setdefault(idx, {})["result"] = (
                    ts, int(m.group(2)), float(m.group(3)))
                continue
    if raw:
        segments.append((name, raw))

    tasks = []
    for seg_name, seg_raw in segments:
        steps = []
        for i in sorted(seg_raw):
            r = seg_raw[i]
            lc = r.get("lc")
            if not lc:
                continue
            lc_ts, pause_s, delay_s, resume_s, retry_s, pause_err = lc
            probe = r.get("probe")
            anchor = probe[0] if probe else lc_ts
            probe_s = (probe[1] + probe[2]) if probe else 0.0
            sched_resume = (r.get("sched") or (0.0, 0.0))[1]
            rc = anchor - timedelta(seconds=probe_s)
            rs = rc - timedelta(seconds=resume_s + retry_s)
            pc = rs - timedelta(seconds=sched_resume + delay_s)
            ps = pc - timedelta(seconds=pause_s)
            res = r.get("result")
            if res:
                end, exit_code, elapsed = res
                start = end - timedelta(seconds=elapsed)
            else:
                start = end = rc
                exit_code, elapsed = None, 0.0
            step = {
                "index": i,
                "action": r.get("action", ""),
                "exit": exit_code,
                "pause_ms": round(pause_s * 1000, 1),
                "paused_ms": round((delay_s + sched_resume) * 1000, 1),
                "resume_ms": round((resume_s + retry_s) * 1000, 1),
                "exec_ms": round(elapsed * 1000, 1),
                "pause_started_at": _iso(ps),
                "pause_completed_at": _iso(pc),
                "resume_started_at": _iso(rs),
                "resume_completed_at": _iso(rc),
                "started_at": _iso(start),
                "completed_at": _iso(end),
            }
            if pause_err != "None":
                step["pause_error"] = pause_err
            steps.append(step)
        tasks.append((seg_name, steps))
    return tasks


V3_TIMING_RE = re.compile(
    r"Agent interaction timing step=(\d+)/\d+ pause=([\d.]+)s "
    r"delay=([\d.]+)s capacity_wait=([\d.]+)s resume=([\d.]+)s "
    r"command=([\d.]+)s total=([\d.]+)s")
V3_BLOCK_CMD_RE = re.compile(r"\[(\S+)\] \[(\d+)/(\d+)\] COMMAND (START|END)")


def parse_runlog_v3(path):
    """Parse current replay-aenv agent logs (main @ec5c928) -> [(name, steps)].

    The agent no longer logs "Replay lifecycle:" or "Lifecycle execution"
    lines. Per step it emits one self-contained timing line:
        Agent interaction timing step=i/N pause=P s delay=D s
        capacity_wait=C s resume=R s command=E s total=T s
    where resume includes the ready probe. In job.log the line is prefixed
    with an ISO timestamp (== command completion); in console/run.log captures
    it is bare, so the anchor falls back to this step's COMMAND END block
    timestamp. Absolute boundaries are reconstructed by subtracting durations:
        command_end      = line ts (or COMMAND END ts)
        resume_completed = command_end - command
        resume_started   = resume_completed - resume
        pause_completed  = resume_started - delay - capacity_wait
        pause_started    = pause_completed - pause
    Action/exit come from job.log "[i/N] COMMAND"/"[i/N] RESULT" lines or the
    console COMMAND START/END blocks. Segments split on "Replay started".
    """
    segments = []
    name = None
    raw = None
    cur = None       # idx whose action lines are being read (console blocks)
    last_end = None  # idx whose exit_code/elapsed lines may follow
    with open(path, errors="replace") as f:
        for line in f:
            ts = None
            msg = line
            if line.startswith("2"):  # ISO-prefixed event line (job.log)
                ts_s, _, m2 = line.partition(" INFO ")
                if m2:
                    ts = datetime.fromisoformat(ts_s.strip())
                    msg = m2
            m = V2_STARTED_RE.search(msg)
            if m:
                if raw:
                    segments.append((name, raw))
                name = os.path.splitext(os.path.basename(m.group(1)))[0]
                raw = {}
                cur = last_end = None
                continue
            if raw is None:
                continue
            m = V3_TIMING_RE.search(msg)
            if m:
                idx = int(m.group(1)) - 1
                e = raw.setdefault(idx, {})
                e["timing"] = tuple(float(g) for g in m.groups()[1:])
                if ts is not None:
                    e["timing_ts"] = ts
                cur = last_end = None
                continue
            c = V3_BLOCK_CMD_RE.search(msg)
            if c:  # console/run.log blocks: "[<ts>] [i/N] COMMAND START|END"
                idx = int(c.group(2)) - 1
                bts = datetime.fromisoformat(c.group(1))
                e = raw.setdefault(idx, {})
                if c.group(4) == "START":
                    e["start_ts"] = bts
                    e["action_lines"] = []
                    cur, last_end = idx, None
                else:
                    e["end_ts"] = bts
                    cur, last_end = None, idx
                continue
            m = V2_CMD_RE.search(msg)
            if m:  # job.log "[i/N] COMMAND <action>"
                idx = int(m.group(1)) - 1
                raw.setdefault(idx, {})["action"] = m.group(2).replace("\\n", " ")
                cur = last_end = None
                continue
            m = V2_RESULT_RE.search(msg)
            if m:  # job.log "[i/N] RESULT exit_code=.. succeeded=.. elapsed=.."
                idx = int(m.group(1)) - 1
                e = raw.setdefault(idx, {})
                e["exit"] = int(m.group(2))
                e["elapsed"] = float(m.group(3))
                if ts is not None:
                    e["end_ts"] = ts
                cur = last_end = None
                continue
            if cur is not None:
                if line.strip():
                    raw[cur]["action_lines"].append(line.rstrip("\n"))
                else:
                    cur = None
                continue
            if last_end is not None:
                if line.startswith("exit_code:"):
                    v = line.split(":", 1)[1].strip()
                    raw[last_end]["exit"] = None if v == "None" else int(v)
                elif line.startswith("elapsed:"):
                    raw[last_end]["elapsed"] = float(
                        line.split(":", 1)[1].strip().rstrip("s"))
                elif line.startswith("stdout:"):
                    last_end = None
                continue
    if raw:
        segments.append((name, raw))

    tasks = []
    for seg_name, seg_raw in segments:
        steps = []
        for i in sorted(seg_raw):
            r = seg_raw[i]
            tm = r.get("timing")
            if not tm:
                continue
            pause_s, delay_s, cap_s, resume_s, cmd_s, _total_s = tm
            end = r.get("timing_ts") or r.get("end_ts")
            if end is None:
                continue
            start = end - timedelta(seconds=cmd_s)
            rc = start  # resume (incl. ready probe) completes as the command starts
            rs = rc - timedelta(seconds=resume_s)
            pc = rs - timedelta(seconds=delay_s + cap_s)
            ps = pc - timedelta(seconds=pause_s)
            action = r.get("action") or " ".join(r.get("action_lines") or [])
            elapsed = r.get("elapsed", cmd_s)
            steps.append({
                "index": i,
                "action": action,
                "exit": r.get("exit"),
                "pause_ms": round(pause_s * 1000, 1),
                "paused_ms": round((delay_s + cap_s) * 1000, 1),
                "resume_ms": round(resume_s * 1000, 1),
                "exec_ms": round(elapsed * 1000, 1),
                "pause_started_at": _iso(ps),
                "pause_completed_at": _iso(pc),
                "resume_started_at": _iso(rs),
                "resume_completed_at": _iso(rc),
                "started_at": _iso(r.get("start_ts", start)),
                "completed_at": _iso(end),
            })
        tasks.append((seg_name, steps))
    return tasks


def parse_runlog_auto(path):
    """Dispatch on content: v3 (Agent interaction timing), v2 (jobs.txt
    Lifecycle execution), or legacy run.log (Replay lifecycle) format.
    Streams until the first marker line — command blocks can push the first
    timing line far past any fixed-size head read."""
    marker = None
    with open(path, errors="replace") as f:
        for line in f:
            if "Agent interaction timing" in line:
                marker = "v3"
                break
            if "Lifecycle execution" in line:
                marker = "v2"
                break
            if "Replay lifecycle:" in line:
                break
    if marker == "v3":
        return parse_runlog_v3(path)
    if marker == "v2":
        return parse_joblog(path)
    return [(None, parse_runlog(path))]


def load_runlog_tasks(runlog_path):
    """Resolve --runlog argument into [(name, steps)].

    - <file>                     -> auto-detect format, single task (new-format
                                    files may contain several trajectories)
    - <dir>/run.log              -> single task (format auto-detected)
    - <dir>/<task*>/run.log      -> multi task (format auto-detected)
    - <dir>/*.txt|*.log (flat)   -> new-format multi, one task per file
    """
    if os.path.isfile(runlog_path):
        stem = os.path.splitext(os.path.basename(runlog_path))[0]
        tasks = [(nm or stem, st) for nm, st in parse_runlog_auto(runlog_path)]
        tasks = [t for t in tasks if t[1]]
        if not tasks:
            sys.exit(f"no lifecycle lines in {runlog_path}")
        return tasks
    direct = os.path.join(runlog_path, "run.log")
    if os.path.isfile(direct):
        dname = os.path.basename(runlog_path.rstrip("/"))
        tasks = [(nm or dname, st) for nm, st in parse_runlog_auto(direct)]
        tasks = [t for t in tasks if t[1]]
        if not tasks:
            sys.exit(f"no lifecycle lines in {direct}")
        return tasks
    tasks = []
    for d in sorted(glob.glob(os.path.join(runlog_path, "*"))):
        p = os.path.join(d, "run.log")
        if not os.path.isfile(p):
            continue
        dname = os.path.basename(d)
        got = False
        for nm, st in parse_runlog_auto(p):
            if not st:
                continue
            tasks.append((nm or dname, st))
            got = True
        if not got:
            print(f"warning: {d} has no lifecycle lines (failed early?), skipped",
                  file=sys.stderr)
    if tasks:
        return tasks
    for p in sorted(glob.glob(os.path.join(runlog_path, "*.txt"))
                    + glob.glob(os.path.join(runlog_path, "*.log"))):
        stem = os.path.splitext(os.path.basename(p))[0]
        for nm, st in parse_runlog_auto(p):
            if not st:
                print(f"warning: {p} has no lifecycle lines, skipped",
                      file=sys.stderr)
                continue
            tasks.append((nm or stem, stem, st))
    if not tasks:
        sys.exit(f"no run.log / new-format log with lifecycle lines under {runlog_path}")
    seen = {}
    out = []
    for nm, stem, st in tasks:
        if nm in seen:  # same trajectory replayed in several files: disambiguate
            nm = f"{nm} [{stem}]"
        seen[nm] = True
        out.append((nm, st))
    return out


def build_windows(steps):
    """Per-step phase windows as (phase, t_start, t_end, step_index)."""
    wins = []
    for s in steps:
        ps, pc = to_epoch(s.get("pause_started_at")), to_epoch(s.get("pause_completed_at"))
        rs, rc = to_epoch(s.get("resume_started_at")), to_epoch(s.get("resume_completed_at"))
        done = to_epoch(s.get("completed_at"))
        i = s.get("index")
        for ph, a, b in (("pause", ps, pc), ("paused", pc, rs),
                         ("resume", rs, rc), ("exec", rc, done)):
            if a is not None and b is not None and b > a:
                wins.append((ph, a, b, i))
    return wins


def phase_of(t, wins):
    for ph, a, b, _ in wins:
        if a <= t < b:
            return ph
    return "gap"


def agg(vals):
    if not vals:
        return None
    return {"n": len(vals), "mean": sum(vals) / len(vals), "max": max(vals)}


def phase_table(per_phase, wins_count, disks):
    L = []
    L.append("| Phase | Samples | CPU% mean/max | iowait% mean/max | FC cores mean/max | Read MB/s mean/max | Write MB/s mean/max | Dirty MB mean/max |")
    L.append("|---|---|---|---|---|---|---|---|")
    label = {"pause": "pause", "paused": "paused (LLM window)",
             "resume": "resume", "exec": "exec", "gap": "gap (between steps)"}
    for ph in PHASES:
        d = per_phase[ph]
        if not d["cpu"]:
            continue
        cells = []
        for k, fmt in (("cpu", "{:.1f}/{:.1f}"), ("iowait", "{:.1f}/{:.1f}"),
                       ("fc_cores", "{:.2f}/{:.2f}"),
                       ("r", "{:.1f}/{:.1f}"), ("w", "{:.1f}/{:.1f}"),
                       ("dirty", "{:.0f}/{:.0f}")):
            vals = d[k]
            if vals:
                cells.append(fmt.format(sum(vals) / len(vals), max(vals)))
            else:
                cells.append("-")
        L.append(f"| {label[ph]} | {len(d['cpu'])} | " + " | ".join(cells) + " |")
    L.append("")
    return L


def assign_phases(metrics, wins, disks):
    """Bucket each metric sample into a phase; return per_phase, sample_ts."""
    per_phase = {p: {"cpu": [], "iowait": [], "fc_cores": [], "r": [], "w": [], "dirty": []}
                 for p in PHASES}
    sample_ts = []
    for m in metrics:
        t = to_epoch(m["ts"])
        sample_ts.append(t)
        ph = phase_of(t, wins)
        d = per_phase[ph]
        d["cpu"].append(m.get("cpu_pct"))
        iw = m.get("cpu_iowait_pct")
        if iw is not None:
            d["iowait"].append(iw)
        d["fc_cores"].append(m.get("fc_cores"))
        dd = m.get("disk") or {}
        d["r"].append(sum((dd[x] or {}).get("r_mb_s", 0) for x in disks))
        d["w"].append(sum((dd[x] or {}).get("w_mb_s", 0) for x in disks))
        d["dirty"].append(m.get("mem_dirty_mb"))
    return per_phase, sample_ts


def cycle_peaks(metrics, sample_ts, wins, disks):
    """Per-cycle peak disk write within pause, peak disk read within resume."""
    cyc_pause_w, cyc_resume_r = [], []
    for ph, a, b, i in wins:
        seg = [m for m, t in zip(metrics, sample_ts) if a <= t < b]
        if not seg:
            continue
        if ph == "pause":
            cyc_pause_w.append(max(
                sum((m["disk"].get(x) or {}).get("w_mb_s", 0) for x in disks)
                for m in seg if m.get("disk")))
        elif ph == "resume":
            cyc_resume_r.append(max(
                sum((m["disk"].get(x) or {}).get("r_mb_s", 0) for x in disks)
                for m in seg if m.get("disk")))
    return cyc_pause_w, cyc_resume_r


def duration_table(tasks):
    """Per-task window durations from log timestamps (no host metrics needed)."""
    L = []
    L.append("## Per-task window durations (from log timestamps, monitor-independent)\n")
    L.append("| Task | cycles | pause ms mean/max | paused ms mean/max "
             "| resume ms mean/max | exec ms mean/max |")
    L.append("|---|---|---|---|---|---|")
    for name, steps in tasks:
        dur = {ph: [] for ph in ("pause", "paused", "resume", "exec")}
        for ph, a, b, _i in build_windows(steps):
            dur[ph].append((b - a) * 1000)
        cells = []
        for ph in ("pause", "paused", "resume", "exec"):
            v = dur[ph]
            cells.append(f"{sum(v)/len(v):.0f}/{max(v):.0f}" if v else "-")
        L.append(f"| {name[:60]} | {len(steps)} | " + " | ".join(cells) + " |")
    L.append("")
    return L


def steps_detail_table(tasks, with_task):
    """Full per-step trajectory table (all steps, all tasks)."""
    L = []
    L.append("## Per-step detail (full trajectory from run.log)\n")
    hdr = ("| Task | " if with_task else "") + \
          "| step | action | pause ms | paused ms | resume ms | exec ms | exit " \
          "| pause start | resume done |"
    sep = ("|---|" if with_task else "") + "|---|---|---|---|---|---|---|---|---|"
    L.append(hdr)
    L.append(sep)
    for name, steps in tasks:
        for s in steps:
            action = (s.get("action") or "").replace("|", "\\|")[:60]
            cells = (
                f"| {s.get('index')} | {action} "
                f"| {s.get('pause_ms','-')} | {s.get('paused_ms','-')} "
                f"| {s.get('resume_ms','-')} | {s.get('exec_ms','-')} "
                f"| {s.get('exit') if s.get('exit') is not None else '-'} "
                f"| {s.get('pause_started_at','')[11:23]} "
                f"| {s.get('resume_completed_at','')[11:23]} |"
            )
            if with_task:
                cells = f"| {name[:40]} " + cells
            L.append(cells)
    L.append("")
    return L


def write_runlog_xlsx(path, tasks, per_phase=None):
    """xlsx full trajectory: per-task summary / per-step detail / phase resources (when monitor data exists)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Per-task summary"
    ws.append(["Task", "cycles", "Wall s", "pause ms mean", "pause ms max",
               "paused ms mean", "paused ms max", "resume ms mean",
               "resume ms max", "exec ms mean", "exec ms max"])
    for name, steps in tasks:
        dur = {ph: [] for ph in ("pause", "paused", "resume", "exec")}
        for ph, a, b, _i in build_windows(steps):
            dur[ph].append((b - a) * 1000)
        t0 = min(to_epoch(s["pause_started_at"]) for s in steps)
        t1 = max(to_epoch(s["completed_at"]) for s in steps)
        row = [name, len(steps), round(t1 - t0, 1)]
        for ph in ("pause", "paused", "resume", "exec"):
            v = dur[ph]
            row += [round(sum(v) / len(v), 1) if v else None,
                    round(max(v), 1) if v else None]
        ws.append(row)

    ws = wb.create_sheet("Per-step detail")
    ws.append(["Task", "step", "action", "pause ms", "paused ms", "resume ms",
               "exec ms", "exit", "pause start", "pause done",
               "resume start", "resume done"])
    for name, steps in tasks:
        for s in steps:
            ws.append([name, s.get("index"), s.get("action"),
                       s.get("pause_ms"), s.get("paused_ms"), s.get("resume_ms"),
                       s.get("exec_ms"), s.get("exit"),
                       s.get("pause_started_at"), s.get("pause_completed_at"),
                       s.get("resume_started_at"), s.get("resume_completed_at")])
    ws.freeze_panes = "A2"

    if per_phase:
        ws = wb.create_sheet("Phase resources")
        ws.append(["Phase", "Samples", "CPU% mean", "CPU% max", "iowait% mean",
                   "iowait% max", "FC cores mean",
                   "FC cores max", "Read MB/s mean", "Read MB/s max", "Write MB/s mean",
                   "Write MB/s max", "Dirty MB mean", "Dirty MB max"])
        for ph in PHASES:
            d = per_phase[ph]
            if not d["cpu"]:
                continue
            row = [ph, len(d["cpu"])]
            for k in ("cpu", "iowait", "fc_cores", "r", "w", "dirty"):
                v = d[k]
                row += [round(sum(v) / len(v), 2) if v else None,
                        round(max(v), 2) if v else None]
            ws.append(row)
        ws.freeze_panes = "A2"

    wb.save(path)
    print(f"wrote {path}")


def analyze_single(obs_dir, label, steps, metrics, plot, runlog=False):
    wins = build_windows(steps)
    if not wins:
        sys.exit("empty windows")
    if not metrics:  # run.log standalone: durations only, no host correlation
        L = [f"# Single-task pause/resume window record (from run.log)\n",
             f"- Task: `{label}`", f"- Steps {len(steps)}\n"]
        L.extend(duration_table([(label, steps)]))
        if runlog:
            L.extend(steps_detail_table([(label, steps)], with_task=False))
            write_runlog_xlsx(os.path.join(obs_dir, "pause_resume_analysis.xlsx"),
                              [(label.replace("runlog:", ""), steps)])
        out_md = os.path.join(obs_dir, "pause_resume_analysis.md")
        with open(out_md, "w") as f:
            f.write("\n".join(L) + "\n")
        print(f"wrote {out_md}")
        print("\n".join(L))
        if plot:
            _plot_multi(obs_dir, [(label, steps)], metrics, [])
        return
    disks = sorted({d for m in metrics for d in (m.get("disk") or {})})
    per_phase, sample_ts = assign_phases(metrics, wins, disks)
    cyc_pause_w, cyc_resume_r = cycle_peaks(metrics, sample_ts, wins, disks)

    L = []
    L.append("# Single-task pause/resume phase host-resource analysis\n")
    L.append(f"- obs: `{obs_dir}`")
    L.append(f"- Task: `{label}`")
    L.append(f"- Sample points {len(metrics)}, steps {len(steps)}, disks {disks}\n")
    L.extend(phase_table(per_phase, len(wins), disks))
    if cyc_pause_w:
        s = sorted(cyc_pause_w)
        L.append(f"Per-cycle pause-window write peak MB/s: mean {sum(s)/len(s):.1f}, "
                 f"p50 {s[len(s)//2]:.1f}, max {s[-1]:.1f} (n={len(s)})")
    if cyc_resume_r:
        s = sorted(cyc_resume_r)
        L.append(f"Per-cycle resume-window read peak MB/s: mean {sum(s)/len(s):.1f}, "
                 f"p50 {s[len(s)//2]:.1f}, max {s[-1]:.1f} (n={len(s)})")
    if runlog:
        L.extend(steps_detail_table([(label, steps)], with_task=False))
        write_runlog_xlsx(os.path.join(obs_dir, "pause_resume_analysis.xlsx"),
                          [(label.replace("runlog:", ""), steps)],
                          per_phase=per_phase)
    out_md = os.path.join(obs_dir, "pause_resume_analysis.md")
    with open(out_md, "w") as f:
        f.write("\n".join(L) + "\n")
    print(f"wrote {out_md}")
    print("\n".join(L))
    if plot:
        _plot(obs_dir, metrics, sample_ts, wins, disks)


def analyze_multi(obs_dir, runlog_dir, tasks, metrics, plot):
    """Multi-task record from a folder of run.log dirs."""
    disks = sorted({d for m in metrics for d in (m.get("disk") or {})})
    all_wins = []
    for name, steps in tasks:
        all_wins.extend(build_windows(steps))

    L = []
    L.append("# Multi-task pause/resume phase host-resource analysis (from run.log)\n")
    L.append(f"- obs: `{obs_dir}`")
    L.append(f"- runlog dir: `{runlog_dir}`")
    L.append(f"- Tasks {len(tasks)}, sample points {len(metrics)}, disks {disks}\n")
    L.extend(duration_table(tasks))
    L.extend(steps_detail_table(tasks, with_task=True))

    if metrics:
        # multi-task phase windows overlap: samples are merged once by priority pause>resume>exec>paused
        pri = ("pause", "resume", "exec", "paused")
        sample_ts = []
        per_phase = {p: {"cpu": [], "iowait": [], "fc_cores": [], "r": [], "w": [], "dirty": []}
                     for p in PHASES}
        for m in metrics:
            t = to_epoch(m["ts"])
            sample_ts.append(t)
            ph = "gap"
            for want in pri:
                if any(p == want and a <= t < b for p, a, b, _ in all_wins):
                    ph = want
                    break
            d = per_phase[ph]
            d["cpu"].append(m.get("cpu_pct"))
            iw = m.get("cpu_iowait_pct")
            if iw is not None:
                d["iowait"].append(iw)
            d["fc_cores"].append(m.get("fc_cores"))
            dd = m.get("disk") or {}
            d["r"].append(sum((dd[x] or {}).get("r_mb_s", 0) for x in disks))
            d["w"].append(sum((dd[x] or {}).get("w_mb_s", 0) for x in disks))
            d["dirty"].append(m.get("mem_dirty_mb"))
        L.append("## Phase host resources (overlapping windows merged as pause>resume>exec>paused)\n")
        L.extend(phase_table(per_phase, len(all_wins), disks))
    else:
        per_phase = None
        L.append("(no monitor data overlapping the log window; phase resource table skipped)\n")

    write_runlog_xlsx(os.path.join(obs_dir, "pause_resume_analysis.xlsx"),
                      tasks, per_phase=per_phase if metrics else None)

    out_md = os.path.join(obs_dir, "pause_resume_analysis_multi.md")
    with open(out_md, "w") as f:
        f.write("\n".join(L) + "\n")
    print(f"wrote {out_md}")
    print("\n".join(L))
    if plot:
        _plot_multi(obs_dir, tasks, metrics, disks)


def main(obs_dir, job_dir, plot=False, runlog=None):
    metrics = load_metrics(obs_dir)
    if runlog:
        os.makedirs(obs_dir, exist_ok=True)
        tasks = load_runlog_tasks(runlog)
        if len(tasks) == 1:
            name, steps = tasks[0]
            analyze_single(obs_dir, f"runlog:{name}", steps, metrics, plot,
                           runlog=True)
        else:
            analyze_multi(obs_dir, runlog, tasks, metrics, plot)
        return
    steps = load_steps(job_dir)
    analyze_single(obs_dir, job_dir, steps, metrics, plot)


def _plot(obs_dir, metrics, sample_ts, wins, disks):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch

    # median-duration full cycle: pause start -> exec end
    cycles = {}
    for ph, a, b, i in wins:
        cycles.setdefault(i, {})[ph] = (a, b)
    full = {i: w for i, w in cycles.items() if len(w) == 4}
    if not full:
        return
    i_med = sorted(full, key=lambda i: full[i]["exec"][1] - full[i]["pause"][0])[len(full) // 2]
    w = full[i_med]
    t0 = w["pause"][0] - 0.5
    t1 = w["exec"][1] + 0.5

    xs, cpu, fcc, rd, wr = [], [], [], [], []
    for m, t in zip(metrics, sample_ts):
        if t0 <= t <= t1:
            xs.append(t - t0)
            cpu.append(m.get("cpu_pct"))
            fcc.append(m.get("fc_cores"))
            dd = m.get("disk") or {}
            rd.append(sum((dd[x] or {}).get("r_mb_s", 0) for x in disks))
            wr.append(sum((dd[x] or {}).get("w_mb_s", 0) for x in disks))

    PCOLOR = {"pause": "#d6272833", "paused": "#ffd8a866",
              "resume": "#9467bd33", "exec": "#2ca02c33"}
    fig, axes = plt.subplots(3, 1, figsize=(14, 10), sharex=True)
    for ax in axes:
        for ph, (a, b) in w.items():
            ax.axvspan(a - t0, b - t0, color=PCOLOR[ph], zorder=0)
    axes[0].plot(xs, wr, label="Write MB/s", color="#d62728")
    axes[0].plot(xs, rd, label="Read MB/s", color="#1f77b4")
    axes[0].set_ylabel("Disk MB/s"); axes[0].legend()
    axes[1].plot(xs, cpu, color="#ff7f0e")
    axes[1].set_ylabel("Host CPU %")
    axes[2].plot(xs, fcc, color="#2ca02c")
    axes[2].set_ylabel("FC CPU cores"); axes[2].set_xlabel(f"t (s), cycle #{i_med} origin = 0.5s before pause")
    axes[0].set_title(f"single pause→resume→exec cycle resource curves (cycle #{i_med}, median duration)")
    for ax in axes:
        ax.grid(linestyle=":", alpha=0.5)
        ax.legend(handles=[Patch(color=PCOLOR[p], label=p) for p in PCOLOR],
                  loc="upper right", fontsize=9)
    plt.tight_layout()
    out = os.path.join(obs_dir, "pause_resume_cycle.png")
    plt.savefig(out, dpi=130)
    plt.close(fig)
    print(f"wrote {out}")


def _plot_multi(obs_dir, tasks, metrics, disks):
    """Per-task phase gantt + host metric curves on a shared time axis."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Patch

    PCOLOR = {"pause": "#d62728", "paused": "#ffd8a866",
              "resume": "#9467bd", "exec": "#2ca02c"}
    spans = []  # (task_row, phase, a, b)
    t_lo = t_hi = None
    for row, (name, steps) in enumerate(tasks):
        for ph, a, b, _i in build_windows(steps):
            spans.append((row, ph, a, b))
            t_lo = a if t_lo is None else min(t_lo, a)
            t_hi = b if t_hi is None else max(t_hi, b)
    if not spans:
        print("no windows, skip multi plot")
        return

    xs, cpu, dirty, wr = [], [], [], []
    for m in metrics:
        t = to_epoch(m["ts"])
        if t_lo <= t <= t_hi:
            xs.append(t - t_lo)
            cpu.append(m.get("cpu_pct"))
            dirty.append(m.get("mem_dirty_mb"))
            dd = m.get("disk") or {}
            wr.append(sum((dd[x] or {}).get("w_mb_s", 0) for x in disks))

    fig, (axg, axm) = plt.subplots(
        2, 1, figsize=(16, 4 + 0.5 * len(tasks)), sharex=True,
        gridspec_kw={"height_ratios": [max(2, len(tasks)), 2]})
    for row, ph, a, b in spans:
        axg.barh(row, b - a, left=a - t_lo, height=0.7,
                 color=PCOLOR[ph], edgecolor="none")
    axg.set_yticks(range(len(tasks)))
    axg.set_yticklabels([n[:40] for n, _ in tasks], fontsize=8)
    axg.invert_yaxis()
    axg.set_ylabel("Task")
    axg.set_title("pause/resume/paused/exec windows (run.log timestamps)")
    axg.legend(handles=[Patch(color=PCOLOR[p], label=p) for p in PCOLOR],
               loc="upper right", fontsize=9)
    axg.grid(axis="x", linestyle=":", alpha=0.5)

    if xs:
        axm.plot(xs, cpu, label="Host CPU %", color="#ff7f0e")
        axm.plot(xs, wr, label="Write MB/s", color="#d62728")
        axm2 = axm.twinx()
        axm2.plot(xs, dirty, label="Dirty MB", color="#8c564b", alpha=0.7)
        axm2.set_ylabel("Dirty MB")
        axm.legend(loc="upper left")
    else:
        axm.text(0.5, 0.5, "monitor data does not overlap the log window", ha="center",
                 transform=axm.transAxes)
    axm.set_ylabel("CPU % / Write MB/s")
    axm.set_xlabel("t (s), relative to first pause start")
    axm.grid(linestyle=":", alpha=0.5)
    plt.tight_layout()
    out = os.path.join(obs_dir, "pause_resume_multi.png")
    plt.savefig(out, dpi=130)
    plt.close(fig)
    print(f"wrote {out}")


if __name__ == "__main__":
    argv = sys.argv[1:]
    plot = "--plot" in argv
    argv = [a for a in argv if a != "--plot"]
    runlog = None
    if "--runlog" in argv:
        i = argv.index("--runlog")
        runlog = argv[i + 1]
        argv = argv[:i] + argv[i + 2:]
    if runlog is not None:
        if len(argv) != 1:
            sys.exit(__doc__)
        main(argv[0], None, plot=plot, runlog=runlog)
    else:
        if len(argv) != 2:
            sys.exit(__doc__)
        main(argv[0], argv[1], plot=plot)
