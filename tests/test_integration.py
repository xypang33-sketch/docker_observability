"""Opt-in Engine test: uses an existing local image with /bin/sh, never pulls.

DOCKER_OBSERVABILITY_TEST_IMAGE=<local-image> python -m unittest discover -s tests -v
"""
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
import uuid

from docker_monitor import Collector, connect
from docker_report import generate


@unittest.skipUnless(os.environ.get("DOCKER_OBSERVABILITY_TEST_IMAGE"), "set DOCKER_OBSERVABILITY_TEST_IMAGE to opt into Docker integration")
class EngineIntegration(unittest.TestCase):
    def test_resource_lifecycle_and_reports(self):
        client = connect()
        run_id = "observability-test-" + uuid.uuid4().hex[:12]
        collector = Collector(client, labels=["observability.run_id=" + run_id])
        container = None
        samples = []
        try:
            # Resolve locally first; no registry download or existing-container mutations.
            image = client.images.get(os.environ["DOCKER_OBSERVABILITY_TEST_IMAGE"])
            container = client.containers.create(
                image.id, name=run_id, entrypoint=["/bin/sh", "-c"],
                command=["while :; do :; done"],
                labels={"observability.run_id": run_id, "observability.task_id": "task-1", "observability.agent_id": "worker-1"},
                nano_cpus=500_000_000, mem_limit="64m",
            )
            container.start()
            samples.append(collector.sample())
            time.sleep(.5)
            samples.append(collector.sample())
            running = samples[-1]["containers"][0]
            self.assertGreater(running["cpu_seconds_delta"], 0)
            self.assertEqual(running["agent_id"], "worker-1")
            self.assertTrue(running["stats_available"])
            container.pause()
            samples.append(collector.sample())
            self.assertEqual(samples[-1]["containers"][0]["status"], "paused")
            container.unpause()
            container.restart(timeout=1)
            samples.append(collector.sample())
            self.assertIsNone(samples[-1]["containers"][0]["cpu_seconds_delta"])
            container.stop(timeout=1)
            samples.append(collector.sample())
            self.assertEqual(samples[-1]["containers"][0]["status"], "exited")
            container.remove()
            container = None
            # Engine removal may return just before its destroy event is queued.
            required = {"create", "start", "pause", "unpause", "die", "destroy"}
            events, actions = [], set()
            deadline = time.monotonic() + 3
            while time.monotonic() < deadline:
                events.extend(collector.events())
                actions = {event["event"] for event in events}
                if required <= actions:
                    break
                time.sleep(.05)
            self.assertTrue(required <= actions, actions)
            self.assertTrue(all(event["run_id"] == run_id for event in events))
            self.assertEqual(collector.events(), [])
            self.assertTrue(all(not s["errors"] for s in samples))
            with tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                (root / "manifest.json").write_text(json.dumps({"backend": "docker", "interval_s": .5}))
                (root / "metrics.jsonl").write_text("\n".join(map(json.dumps, samples)))
                (root / "container_events.jsonl").write_text("\n".join(map(json.dumps, events)))
                generate(root, xlsx=True)
                from openpyxl import load_workbook
                wb = load_workbook(root / "report.xlsx")
                self.assertIn("Concurrent totals", wb.sheetnames)
                self.assertEqual(wb["Containers"].max_row, 2)
                wb.close()
        finally:
            if container is not None:
                container.remove(force=True)
            client.close()
