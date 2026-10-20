"""Generate a Docker resource report without replay/Harbor logs."""

import argparse
from collections import defaultdict
import json
from pathlib import Path


def read_jsonl(path):
    with path.open() as stream:
        return [json.loads(line) for line in stream if line.strip()]


def summarize(samples):
    rows = {}
    for sample in samples:
        for item in sample["containers"]:
            cid = item["container_id"]
            row = rows.setdefault(cid, {
                "container_id": cid, "name": None, "task_id": None, "agent_id": None,
                "run_id": None, "samples": 0, "stats_samples": 0, "errors": 0,
                "observed_cpu_seconds": None, "covered_cpu_interval_s": 0.0,
                "peak_memory_usage_bytes": None, "peak_working_set_bytes": None,
                "peak_cpu_cores": None, "peak_pids": None,
            })
            row["samples"] += 1
            row["errors"] += bool(item.get("error"))
            row["stats_samples"] += bool(item.get("stats_available"))
            for key in ("name", "task_id", "agent_id", "run_id"):
                if item.get(key) is not None:
                    row[key] = item[key]
            delta, span = item.get("cpu_seconds_delta"), item.get("cpu_interval_s")
            if delta is not None and span is not None:
                row["observed_cpu_seconds"] = (row["observed_cpu_seconds"] or 0) + delta
                row["covered_cpu_interval_s"] += span
            for src, dest in (("memory_usage_bytes", "peak_memory_usage_bytes"),
                              ("memory_working_set_bytes", "peak_working_set_bytes"),
                              ("cpu_cores", "peak_cpu_cores"), ("pids", "peak_pids")):
                if item.get(src) is not None:
                    row[dest] = max(row[dest] or 0, item[src])
    for row in rows.values():
        span = row["covered_cpu_interval_s"]
        row["mean_cpu_cores"] = row["observed_cpu_seconds"] / span if span else None
    return list(rows.values())


def show(value):
    if value is None:
        return "unknown"
    if isinstance(value, float):
        return f"{value:.3f}"
    return str(value).replace("|", "\\|").replace("\n", " ")


def collection_status(samples, rows):
    if not samples:
        return "无资源样本：采集未完成或未启动。"
    if not rows:
        if any(s.get("list_ok") for s in samples):
            return "未匹配到任何容器，无法生成资源统计。请检查容器名、label 或 Compose 项目；这不代表 CPU/内存使用量为零。"
        return "容器列表查询失败，无法生成资源统计。请查看采集错误。"
    if not any(row["stats_samples"] for row in rows):
        return "已匹配容器，但没有有效资源数据。请检查容器是否运行以及 stats 采集错误。"
    return "已采集到容器资源数据。"


def selection_text(manifest):
    return json.dumps({key: manifest.get(key) for key in ("containers", "labels", "compose_project")}, ensure_ascii=False)


def markdown(manifest, samples, events, rows):
    failures = sum(not s.get("list_ok") for s in samples)
    errors = sum(len(s.get("errors") or []) for s in samples)
    lines = [
        "# Docker 资源报告", "",
        f"**采集状态：{collection_status(samples, rows)}**", "",
        f"筛选条件：`{show(selection_text(manifest))}`。", "",
        f"采集开始：{manifest.get('started_at')}；结束：{manifest.get('stopped_at', '仍在采集')}。", "",
        f"Engine：{manifest.get('engine_os')}；逻辑 CPU：{manifest.get('engine_cpus')}；采样间隔：{manifest.get('interval_s')} 秒。", "",
        f"样本 {len(samples)}；容器 {len(rows)}；列表查询失败 {failures}；采集错误 {errors}。", "",
        "范围：Docker Engine 中选定的 Linux 容器。Docker Desktop 下属于 Linux VM；不包含 Mac 应用、宿主编排器或远端模型计算。", "",
        "CPU 秒仅合计同一次容器运行中、连续有效计数之间的增量；首个样本之前和最后样本之后的用量不计。unknown 表示缺失，不等于零。", "",
        "内存峰值是采样观测峰值。usage 包含缓存；working set 扣除 inactive file cache，不是 RSS/PSS。容器独立峰值不能相加当作同时峰值。", "",
        "| 容器 | task | agent | 样本/有效stats | CPU秒（已覆盖） | CPU覆盖秒 | 平均核数 | 峰值核数 | 内存峰值MiB | 工作集峰值MiB | 错误 |",
        "|---|---|---|---|---|---|---|---|---|---|---|",
    ]
    for row in rows:
        vals = [row['name'] or row['container_id'][:12], row['task_id'], row['agent_id'],
                f"{row['samples']}/{row['stats_samples']}", row['observed_cpu_seconds'], row['covered_cpu_interval_s'],
                row['mean_cpu_cores'], row['peak_cpu_cores'],
                row['peak_memory_usage_bytes'] / 2**20 if row['peak_memory_usage_bytes'] is not None else None,
                row['peak_working_set_bytes'] / 2**20 if row['peak_working_set_bytes'] is not None else None, row['errors']]
        lines.append("| " + " | ".join(show(v) for v in vals) + " |")
    counts = defaultdict(int)
    for event in events:
        if event.get("container_id"):
            counts[event["event"]] += 1
    lines += ["", "## 容器事件", "", ", ".join(f"{show(k)}={v}" for k, v in sorted(counts.items())) or "无容器事件。", "",
              "事件来自 Docker daemon；采集前已经存在的容器会出现在资源样本中，不补造 create 事件。daemon 的事件缓冲有限，长时间断连或大量事件可能丢失。", "",
              "## 原始数据", "", "- `metrics.jsonl`：每容器资源计数、时间窗口、错误。",
              "- `container_events.jsonl`：create/start/pause/unpause/die/destroy 等事件。",
              "- `summary.json`：按容器汇总，保留 run/task/agent 标签。", "",
              "本报告不推断任务成功率、协作收益或模型等待时长；这些需要关联应用轨迹。", ""]
    return "\n".join(lines)


def spreadsheet(path, rows, samples, events, manifest):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, records):
        ws = wb.create_sheet(name)
        if not records:
            ws.append(["No observations"])
            return ws
        keys = list(dict.fromkeys(key for row in records for key in row))
        ws.append(keys)
        for row in records:
            values = []
            for key in keys:
                value = row.get(key)
                if isinstance(value, (list, dict)):
                    value = json.dumps(value, ensure_ascii=False)
                values.append(value)
            ws.append(values)
            # Docker labels/names are untrusted strings, not Excel formulas.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    sheet("Scope and units", [
        {"item": "采集状态", "value": collection_status(samples, rows)},
        {"item": "筛选条件", "value": selection_text(manifest)},
        {"item": "scope", "value": "Selected Linux containers on Docker Engine; Desktop means Linux VM, not the Mac host"},
        {"item": "engine_os", "value": manifest.get("engine_os")},
        {"item": "interval_s", "value": manifest.get("interval_s")},
        {"item": "CPU seconds", "value": "Sum of observed counter deltas in the same container start; excludes unsampled startup/exit tails"},
        {"item": "CPU cores", "value": "CPU seconds / actual counter interval; one means one logical CPU"},
        {"item": "Memory bytes", "value": "usage includes cache; working set excludes inactive file cache, not RSS/PSS; peaks are sampled"},
        {"item": "Missing values", "value": "Blank is unavailable, not zero; see Errors and per-container error fields"},
        {"item": "Concurrent totals", "value": "Sum within a sampling pass; container reads are near-concurrent, not an atomic snapshot"},
        {"item": "Quality", "value": "Task success, model latency and collaboration benefits require application traces"},
    ])
    if rows:
        sheet("Containers", rows)
    else:
        sheet("Containers", [{"诊断": collection_status(samples, rows), "筛选条件": selection_text(manifest)}])
    sheet("Samples", [{"sample_ts": s["ts"], **c} for s in samples for c in s["containers"]])
    sheet("Events", events)
    sheet("Errors", [{"ts": s["ts"], **error} for s in samples for error in s.get("errors", [])])
    totals = []
    for sample in samples:
        running = [c for c in sample["containers"] if c.get("status") in ("running", "paused")]
        valid = [c for c in running if c.get("stats_available")]
        values = {}
        for key in ("cpu_cores", "memory_usage_bytes", "memory_working_set_bytes"):
            values[key] = sum(c[key] for c in valid) if sample.get("containers") and sample.get("list_ok") and len(valid) == len(running) and all(c.get(key) is not None for c in valid) and not any(c.get("error") for c in sample["containers"]) else None
        totals.append({"ts": sample["ts"], **values})
    ws = sheet("Concurrent totals", totals)
    if totals:
        chart = LineChart()
        chart.title = "Selected containers: observed CPU cores"
        chart.y_axis.title = "CPU cores"
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        ws.add_chart(chart, "F2")
    wb.save(path)


def generate(obs_dir, *, xlsx=False):
    obs_dir = Path(obs_dir)
    manifest = json.loads((obs_dir / "manifest.json").read_text())
    if manifest.get("backend") != "docker":
        raise ValueError("expected Docker observations; AgentEnv tools are under legacy/agentenv")
    samples = read_jsonl(obs_dir / "metrics.jsonl")
    events = read_jsonl(obs_dir / "container_events.jsonl")
    rows = summarize(samples)
    (obs_dir / "summary.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n")
    (obs_dir / "report.md").write_text(markdown(manifest, samples, events, rows))
    if xlsx:
        spreadsheet(obs_dir / "report.xlsx", rows, samples, events, manifest)
    return rows


def cli(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("obs_dir", type=Path)
    parser.add_argument("--xlsx", action="store_true")
    args = parser.parse_args(argv)
    generate(args.obs_dir, xlsx=args.xlsx)
    print(f"wrote report to {args.obs_dir}")


if __name__ == "__main__":
    cli()
